#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
MEXC Futures — радар ближних к недельному LLV/HHV + расширенные фильтры, Month1-проверка и Excel с подсветкой

Новое в этой версии (v2):
  • Проверка месячных лоёв (Month1) для тикеров, прошедших недельный фильтр
  • Бэйджи в отчёте: 🧱 (LLV26≈LLV52), 📈 (trend_ok), 💰 (wk_amount_ratio≥порога), ✓WM (месячный и недельный лои рядом)
  • Новый лист Near_LL (реально близко к недельному лою) + Legend и Settings
  • (Опционально) жёсткий фильтр по близости к недельному лою: --max-dist-w-pct N (0=выкл)
  • Безопасный CSV при пустом наборе

Зависимости:
    pip install requests aiohttp pandas openpyxl
Опционально:
    pip install gspread google-auth

Python 3.9+
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import math
import os
import sys
import time
from datetime import datetime, timezone
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

import requests

BASE = "https://contract.mexc.com"
WEEK_SECS = 7 * 24 * 3600
MONTH_SECS = 30 * 24 * 3600  # усреднённо

# ---------------------------- утилиты ----------------------------

def log(*a: object) -> None:
    print(*a, flush=True)


def jget(path: str, params: Optional[Dict] = None, timeout: int = 25):
    """Синхронный GET; если у ответа есть поле data — возвращаем его."""
    url = f"{BASE}{path}"
    r = requests.get(url, params=params, timeout=timeout)
    r.raise_for_status()
    js = r.json()
    if isinstance(js, dict) and ("success" in js or "code" in js):
        return js.get("data")
    return js

# ----------------------- market data (sync) ----------------------

def get_all_tickers() -> Dict[str, Dict]:
    data = jget("/api/v1/contract/ticker")
    out: Dict[str, Dict] = {}
    rows = data if isinstance(data, list) else ([data] if isinstance(data, dict) else [])
    for row in rows:
        sym = row.get("symbol")
        if not sym:
            continue

        def f(key: str, default=None):
            v = row.get(key)
            try:
                return float(v) if v is not None else default
            except Exception:
                return default

        out[sym] = {
            "lastPrice": f("lastPrice", f("fairPrice", f("indexPrice"))),
            "bid1": f("bid1", None),
            "ask1": f("ask1", None),
            "volume24": f("volume24", 0.0),
            "amount24": f("amount24", 0.0),
            "holdVol": f("holdVol", 0.0),
            "indexPrice": f("indexPrice"),
            "fairPrice": f("fairPrice"),
        }
    return out


def get_contracts(quote_filter: str = "USDT") -> List[Dict]:
    data = jget("/api/v1/contract/detail")
    items: List[Dict] = []
    if isinstance(data, list):
        items = data
    elif isinstance(data, dict) and data:
        items = [data]

    out: List[Dict] = []
    if items:
        for it in items:
            try:
                st = it.get("state", it.get("status", 0))
                if str(st) not in ("0", "NORMAL", "normal", "ENABLED", "enabled", "1"):
                    st_str = str(st).lower()
                    if any(k in st_str for k in ("stop", "delist", "suspend")):
                        continue
                if quote_filter:
                    want = quote_filter.upper()
                    q = str(it.get("quoteCoin", "")).upper()
                    sym = str(it.get("symbol", "")).upper()
                    if (q and q != want) and (sym and not sym.endswith(f"_{want}")):
                        continue
                out.append(it)
            except Exception:
                continue
        if out:
            return out

    # Фолбэк: из тикеров
    ticks = get_all_tickers()
    if not ticks:
        return []
    for sym in sorted(ticks.keys()):
        if quote_filter and not sym.upper().endswith(f"_{quote_filter.upper()}"):
            continue
        out.append({"symbol": sym})
    return out

# --------------- асинхронная загрузка kline (Week/Month) ---------

async def fetch_klines(session, symbol: str, weeks: int, retries: int = 3) -> Tuple[str, Optional[Dict]]:
    import aiohttp

    end_ts = int(time.time())
    start_ts = end_ts - (weeks + 20) * WEEK_SECS
    params = {"interval": "Week1", "start": start_ts, "end": end_ts}
    url = f"{BASE}/api/v1/contract/kline/{symbol}"
    backoff = 0.6
    for _ in range(retries):
        try:
            async with session.get(url, params=params, timeout=aiohttp.ClientTimeout(total=20)) as resp:
                if resp.status == 200:
                    js = await resp.json()
                    data = js.get("data") if isinstance(js, dict) and ("success" in js or "code" in js) else js
                    if isinstance(data, dict) and data.get("time"):
                        return symbol, data
                    return symbol, None
                if resp.status in (429, 500, 502, 503, 504):
                    await asyncio.sleep(backoff)
                    backoff *= 1.7
                    continue
                return symbol, None
        except Exception:
            await asyncio.sleep(backoff)
            backoff *= 1.7
    return symbol, None


async def fetch_klines_month(session, symbol: str, months: int, retries: int = 3) -> Tuple[str, Optional[Dict]]:
    import aiohttp

    end_ts = int(time.time())
    start_ts = end_ts - (months + 14) * MONTH_SECS
    params = {"interval": "Month1", "start": start_ts, "end": end_ts}
    url = f"{BASE}/api/v1/contract/kline/{symbol}"
    backoff = 0.6
    for _ in range(retries):
        try:
            async with session.get(url, params=params, timeout=aiohttp.ClientTimeout(total=20)) as resp:
                if resp.status == 200:
                    js = await resp.json()
                    data = js.get("data") if isinstance(js, dict) and ("success" in js or "code" in js) else js
                    if isinstance(data, dict) and data.get("time"):
                        return symbol, data
                    return symbol, None
                if resp.status in (429, 500, 502, 503, 504):
                    await asyncio.sleep(backoff)
                    backoff *= 1.7
                    continue
                return symbol, None
        except Exception:
            await asyncio.sleep(backoff)
            backoff *= 1.7
    return symbol, None


async def gather_klines(symbols: Sequence[str], weeks: int, concurrency: int = 10) -> Dict[str, Dict]:
    try:
        import aiohttp  # noqa: F401
    except Exception:
        return {}

    import aiohttp

    sem = asyncio.Semaphore(max(1, concurrency))
    out: Dict[str, Dict] = {}

    async def worker(sym: str) -> None:
        async with sem:
            async with aiohttp.ClientSession() as session:
                s, data = await fetch_klines(session, sym, weeks)
                if data:
                    out[s] = data

    tasks = [asyncio.create_task(worker(s)) for s in symbols]
    done = 0
    for task in asyncio.as_completed(tasks):
        await task
        done += 1
        if done % 20 == 0:
            log(f"…kline Week1 получены: {done}/{len(symbols)}")
    return out


async def gather_klines_month(symbols: Sequence[str], months: int, concurrency: int = 8) -> Dict[str, Dict]:
    try:
        import aiohttp  # noqa: F401
    except Exception:
        return {}

    import aiohttp

    sem = asyncio.Semaphore(max(1, concurrency))
    out: Dict[str, Dict] = {}

    async def worker(sym: str) -> None:
        async with sem:
            async with aiohttp.ClientSession() as session:
                s, data = await fetch_klines_month(session, sym, months)
                if data:
                    out[s] = data

    tasks = [asyncio.create_task(worker(s)) for s in symbols]
    done = 0
    for task in asyncio.as_completed(tasks):
        await task
        done += 1
        if done % 20 == 0:
            log(f"…kline Month1 получены: {done}/{len(symbols)}")
    return out

# ------------------------- математика/метрики ---------------------

def compute_hhv_llv(highs: Sequence[float], lows: Sequence[float], lookback: int) -> Tuple[float, float]:
    if not highs or not lows:
        return (math.nan, math.nan)
    n = min(lookback, len(highs), len(lows))
    return max(highs[-n:]), min(lows[-n:])


def pct_in_range(price: float, lo: float, hi: float) -> float:
    if any(map(math.isnan, (price, lo, hi))) or hi <= lo:
        return math.nan
    return 100.0 * (price - lo) / (hi - lo)


def normalize_log(values: Sequence[float]) -> Dict[int, float]:
    logs: List[Optional[float]] = []
    idxs: List[int] = []
    for i, v in enumerate(values):
        if v is None or v <= 0:
            logs.append(None)
            idxs.append(i)
            continue
        logs.append(math.log(v + 1.0))
        idxs.append(i)
    valid = [x for x in logs if x is not None]
    if not valid:
        return {i: 0.0 for i in idxs}
    lo, hi = min(valid), max(valid)
    rng = (hi - lo) if hi > lo else 1e-9
    return {i: (0.0 if logs[i] is None else (logs[i] - lo) / rng) for i in idxs}


def last_complete_week_index(times: Sequence[int], now_ts: int) -> Optional[int]:
    if not times:
        return None
    if now_ts - times[-1] >= WEEK_SECS:
        return -1
    return -2 if len(times) >= 2 else None

# === helpers ===

def ema_series(values: Sequence[float], period: int) -> List[float]:
    vals = list(values)
    if period <= 1 or len(vals) == 0:
        return vals[:]
    k = 2.0 / (period + 1.0)
    out: List[float] = []
    ema = vals[0]
    out.append(ema)
    for v in vals[1:]:
        ema = v * k + ema * (1.0 - k)
        out.append(ema)
    return out


def atr_sma(highs: Sequence[float], lows: Sequence[float], closes: Sequence[float], period: int, upto_idx: int) -> Optional[float]:
    n = upto_idx + 1
    if n < 2:
        return None
    tr: List[float] = []
    for i in range(1, n):
        h = highs[i]
        l = lows[i]
        pc = closes[i - 1]
        tr.append(max(h - l, abs(h - pc), abs(l - pc)))
    if len(tr) < period:
        return None
    return sum(tr[-period:]) / float(period)


def percentile_rank(values: Sequence[float], v: float) -> float:
    if not values:
        return 0.0
    arr = sorted([x for x in values if isinstance(x, (int, float)) and not math.isnan(x)])
    if not arr:
        return 0.0
    import bisect

    pos = bisect.bisect_right(arr, v)
    return pos / float(len(arr))


def corr_pearson(a: Sequence[float], b: Sequence[float]) -> Optional[float]:
    n = min(len(a), len(b))
    if n < 3:
        return None
    ax = list(a[-n:])
    bx = list(b[-n:])
    ma = sum(ax) / n
    mb = sum(bx) / n
    num = sum((x - ma) * (y - mb) for x, y in zip(ax, bx))
    da = math.sqrt(sum((x - ma) ** 2 for x in ax))
    db = math.sqrt(sum((y - mb) ** 2 for y in bx))
    if da == 0 or db == 0:
        return None
    return num / (da * db)

# ---------------------------- отчёты/экспорт ----------------------

def send_telegram(text: str) -> bool:
    token = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
    chat_id = os.getenv("TELEGRAM_CHAT_ID", "").strip()
    if not token or not chat_id:
        return False
    url = f"https://api.telegram.org/bot{token}/sendMessage"
    try:
        r = requests.post(
            url,
            json={"chat_id": chat_id, "text": text, "parse_mode": "HTML"},
            timeout=20,
        )
        return r.ok
    except Exception:
        return False


def to_dataframe(rows: Sequence[Dict]):
    try:
        import pandas as pd

        return pd.DataFrame(rows)
    except Exception:
        return None


def save_xlsx_multi(sheets: Dict[str, List[Dict]], path: str) -> bool:
    try:
        import pandas as pd
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import PatternFill
        from openpyxl.utils import get_column_letter
    except Exception:
        return False

    try:
        with pd.ExcelWriter(path, engine="openpyxl") as xw:
            for name, rows in sheets.items():
                if not rows:
                    pd.DataFrame({"empty": []}).to_excel(xw, sheet_name=name, index=False)
                else:
                    pd.DataFrame(rows).to_excel(xw, sheet_name=name, index=False)
            wb = xw.book
            green_wm = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            green_tr = PatternFill(start_color="A9D08E", end_color="A9D08E", fill_type="solid")
            yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            blue = PatternFill(start_color="CFE2F3", end_color="CFE2F3", fill_type="solid")
            nearfill = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
            for name in wb.sheetnames:
                ws = wb[name]
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                headers = [cell.value for cell in ws[1]]
                if not headers:
                    continue
                last_col_letter = get_column_letter(ws.max_column)
                last_row = ws.max_row
                if last_row < 2:
                    continue
                rng = f"A2:{last_col_letter}{last_row}"
                if "wm_llv_close" in headers:
                    wm_col_letter = get_column_letter(headers.index("wm_llv_close") + 1)
                    ws.conditional_formatting.add(
                        rng,
                        FormulaRule(formula=[f"=${wm_col_letter}2=TRUE"], fill=green_wm),
                    )
                if "status" in headers:
                    st_col_letter = get_column_letter(headers.index("status") + 1)
                    ws.conditional_formatting.add(
                        rng,
                        FormulaRule(formula=[f"=${st_col_letter}2=\"TRIGGER\""], fill=green_tr),
                    )
                    ws.conditional_formatting.add(
                        rng,
                        FormulaRule(formula=[f"=${st_col_letter}2=\"READY\""], fill=yellow),
                    )
                    ws.conditional_formatting.add(
                        rng,
                        FormulaRule(formula=[f"=${st_col_letter}2=\"WATCH\""], fill=blue),
                    )
                if name == "Near_LL":
                    for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=ws.max_column):
                        for cell in row:
                            cell.fill = nearfill
        return True
    except Exception:
        return False


def upload_to_gsheet(rows: Sequence[Dict], title: str, worksheet_name: Optional[str] = None) -> bool:
    try:
        import gspread
        from google.oauth2.service_account import Credentials

        creds_path = os.environ["GOOGLE_APPLICATION_CREDENTIALS"]
        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]
        creds = Credentials.from_service_account_file(creds_path, scopes=scopes)
        gc = gspread.authorize(creds)
        try:
            sh = gc.open(title)
        except gspread.SpreadsheetNotFound:
            sh = gc.create(title)
        ws = sh.add_worksheet(
            worksheet_name or datetime.utcnow().strftime("Scan %Y-%m-%d %H:%M"),
            rows=2,
            cols=2,
        )
        if rows:
            header = list(rows[0].keys())
            values = [header] + [[str(r.get(k, "")) for k in header] for r in rows]
            ws.update("A1", values)
        return True
    except Exception:
        return False

# ------------------------- чёрный список --------------------------

def parse_blacklist(blacklist_csv: str, blacklist_file: str) -> Set[str]:
    bl: Set[str] = set()
    if blacklist_csv:
        for tok in blacklist_csv.replace(";", ",").split(","):
            tok = tok.strip().upper()
            if tok:
                bl.add(tok)
    if blacklist_file and os.path.isfile(blacklist_file):
        with open(blacklist_file, "r", encoding="utf-8") as f:
            for line in f:
                for tok in line.replace(";", ",").split(","):
                    tok = tok.strip().upper()
                    if tok:
                        bl.add(tok)
    return bl

# ----------------------------- main ------------------------------

def _safe_float(v: object) -> Optional[float]:
    if isinstance(v, (int, float)):
        if math.isnan(v):
            return None
        return float(v)
    return None


def _is_number(v: object) -> bool:
    return isinstance(v, (int, float)) and not math.isnan(v)

def main() -> None:
    ap = argparse.ArgumentParser(
        description=(
            "MEXC Futures: недельные HHV/LLV, близость к лою, ликвидность/объёмы/спред + Month1 и цветной Excel"
        )
    )
    ap.add_argument("--lookback-weeks", type=int, default=52, help="Окно HHV/LLV (недели)")
    ap.add_argument("--quote", type=str, default="USDT", help="Котируемая валюта (например USDT; пусто = без фильтра)")
    # Ликвидность (по тикерам за 24ч)
    ap.add_argument("--min-amount24", type=float, default=300_000, help="Мин. оборот amount24 (USDT) за 24ч")
    ap.add_argument("--min-volume24", type=float, default=60_000, help="Мин. объём volume24 (контракты) за 24ч")
    ap.add_argument("--min-holdvol", type=float, default=0, help="Мин. OI holdVol (контракты)")
    ap.add_argument("--exclude-illiquid", action="store_true", help="Не трогать неликвиды (исключить из выдачи)")
    # История/возраст
    ap.add_argument("--min-history-weeks", type=int, default=26, help="Мин. число недельных баров в истории (0=выкл)")
    ap.add_argument("--min-age-weeks", type=int, default=8, help="Исключить листинги моложе N недель (0=выкл)")
    # Объёмы по последней завершённой неделе (из /kline)
    ap.add_argument("--min-weekly-amount", type=float, default=0, help="Мин. amount за последнюю завершённую неделю (USDT)")
    ap.add_argument("--min-weekly-vol", type=float, default=0, help="Мин. vol за последнюю завершённую неделю (контракты)")
    # Спред (из /ticker: bid1/ask1)
    ap.add_argument("--max-spread-bps", type=float, default=0, help="Макс. спред в бипсах (0=выкл). Пример: 10 = 0.10%%")
    # Чёрный список
    ap.add_argument("--blacklist", type=str, default="", help="CSV тикеров для исключения (e.g. BTC_USDT,ETH_USDT)")
    ap.add_argument("--blacklist-file", type=str, default="", help="Путь к файлу с чёрным списком")
    # Score и сортировка (для консоли)
    ap.add_argument("--w-prox", type=float, default=0.7, help="Вес близости к LLV (0..1)")
    ap.add_argument("--w-liq", type=float, default=0.3, help="Вес ликвидности (0..1)")
    ap.add_argument(
        "--sort-by",
        type=str,
        default="score",
        choices=["score", "pos", "amount24", "dist_w"],
        help="Ключ сортировки для консоли",
    )
    # Отчёты
    ap.add_argument("--top", type=int, default=25, help="Сколько лучших показать в консоли")
    ap.add_argument("--csv", type=str, default="mexc_weekly_llv_scan.csv", help="CSV файл")
    ap.add_argument("--xlsx", type=str, default="mexc_weekly_llv_scan.xlsx", help="XLSX файл")
    ap.add_argument("--gsheet", type=str, default="", help="Название Google Sheets")
    ap.add_argument("--telegram", action="store_true", help="Отправить Telegram-репорт по топ-N")
    # Async
    ap.add_argument("--concurrency", type=int, default=10, help="Одновременных запросов kline (aiohttp)")
    # Month1
    ap.add_argument("--month-check", action="store_true", help="Подгрузить Month1 и добавить LLV по месяцам")
    ap.add_argument("--lookback-months", type=int, default=12, help="Окно HHV/LLV (месяцы)")
    ap.add_argument("--wm-llv-close-bps", type=float, default=100.0, help="Порог близости недельного и месячного лоёв в бипсах (100=1%%)")
    # Расширенные метрики
    ap.add_argument("--lookback-weeks2", type=int, default=26, help="Окно LLV (недели) для быстрой оценки (LLV26)")
    ap.add_argument("--llv26-close-bps", type=float, default=150.0, help="Порог близости LLV26 к LLV52 (б.п.)")
    ap.add_argument("--atrw-period", type=int, default=14, help="ATR период по неделям")
    ap.add_argument("--atrw-pct-max", type=float, default=6.0, help="Макс. ATR%% для READY/TRIGGER")
    ap.add_argument("--wk-amount-min-ratio", type=float, default=1.2, help="Мин. отношение weekly_amount к медиане 26w")
    ap.add_argument("--touch-weeks", type=int, default=6, help="Сколько недель смотреть для касаний уровня")
    ap.add_argument("--touch-eps-bps", type=float, default=75.0, help="Эпсилон касания уровня в б.п. от LLV52")
    ap.add_argument("--ema-guard", action="store_true", help="Требовать close≥EMA30≥EMA50 для TRIGGER")
    ap.add_argument("--near-ll-pct", type=float, default=1.5, help="Порог близости к LLV52 (в %%) для статуса WATCH")
    ap.add_argument("--watch-spread-bps", type=float, default=30.0, help="Макс. спред (б.п.) для WATCH")
    ap.add_argument(
        "--max-dist-w-pct",
        type=float,
        default=0.0,
        help="Жёсткий фильтр: исключить пары с dist_from_llv_pct > N%% (0=выкл)",
    )
    ap.add_argument("--btc-corr", action="store_true", help="Считать корреляцию к BTC_USDT за 26 недель")
    ap.add_argument("--btc-corr-max", type=float, default=0.7, help="Макс. допустимая корреляция к BTC для отбора")
    ap.add_argument("--stop-max-pct", type=float, default=5.0, help="Макс. расстояние до стопа (LLV52) в %%)")
    ap.add_argument("--score2", action="store_true", help="Считать композитный score_v2 и статусы WATCH/READY/TRIGGER")
    ap.add_argument("--w2-near", type=float, default=0.4, help="Вес близости к LLV52 в score_v2")
    ap.add_argument("--w2-liq", type=float, default=0.2, help="Вес ликвидности в score_v2")
    ap.add_argument("--w2-vol", type=float, default=0.2, help="Вес объёмного подтверждения в score_v2")
    ap.add_argument("--w2-trend", type=float, default=0.2, help="Вес тренда/EMA в score_v2")

    args = ap.parse_args()

    # нормировка весов для старого score
    s = max(1e-9, args.w_prox + args.w_liq)
    w_prox = args.w_prox / s
    w_liq = args.w_liq / s

    need_weeks = max(args.lookback_weeks, args.min_history_weeks, args.min_age_weeks, 1)

    log("→ Контракты…")
    contracts = get_contracts(quote_filter=args.quote) if args.quote is not None else get_contracts("")
    symbols_all = [c.get("symbol") for c in contracts if c.get("symbol")]
    if not symbols_all:
        log("Контракты не найдены. Проверьте сеть/VPN/фаервол или попробуйте --quote \"\".")
        sys.exit(2)
    log(f"✔ {len(symbols_all)} контрактов ({args.quote or 'ANY'})")

    blacklist = parse_blacklist(args.blacklist, args.blacklist_file)
    if blacklist:
        before = len(symbols_all)
        symbols_all = [s for s in symbols_all if s.upper() not in blacklist]
        log(f"→ Чёрный список: исключено {before - len(symbols_all)}")

    log("→ Тикеры (/contract/ticker)…")
    tickers = get_all_tickers()
    if not tickers:
        log("⚠ Тикеры не пришли — будем использовать последний close из свечи.")

    liq: Dict[str, Dict] = {}
    for sym in symbols_all:
        t = tickers.get(sym, {})
        liq[sym] = {
            "amount24": float(t.get("amount24") or 0.0),
            "volume24": float(t.get("volume24") or 0.0),
            "holdVol": float(t.get("holdVol") or 0.0),
            "lastPrice": float(t.get("lastPrice") or 0.0),
            "bid1": (None if t.get("bid1") is None else float(t.get("bid1"))),
            "ask1": (None if t.get("ask1") is None else float(t.get("ask1"))),
        }

    log(f"→ Klines Week1 (async), concurrency={args.concurrency} …")
    try:
        kl_map = asyncio.run(gather_klines(symbols_all, weeks=need_weeks, concurrency=args.concurrency))
    except Exception:
        kl_map = {}

    if not kl_map:
        log("⚠ async недоступен — перехожу на синхронную загрузку kline.")
        kl_map = {}
        for i, sym in enumerate(symbols_all, 1):
            try:
                end_ts = int(time.time())
                start_ts = end_ts - (need_weeks + 20) * WEEK_SECS
                data = jget(
                    f"/api/v1/contract/kline/{sym}",
                    params={"interval": "Week1", "start": start_ts, "end": end_ts},
                )
                if isinstance(data, dict) and data.get("time"):
                    kl_map[sym] = data
            except Exception:
                pass
            if i % 20 == 0:
                log(f"…получено {i}/{len(symbols_all)}")

    now_ts = int(time.time())
    rows: List[Dict] = []

    for sym in symbols_all:
        kl = kl_map.get(sym)
        if not kl:
            continue
        try:
            highs = [float(x) for x in kl.get("high", [])]
            lows = [float(x) for x in kl.get("low", [])]
            closes = [float(x) for x in kl.get("close", [])]
            opens = [float(x) for x in kl.get("open", [])] if "open" in kl else []
            times = [int(x) for x in kl.get("time", [])]
            amnts = [float(x) for x in kl.get("amount", [])] if "amount" in kl else []
            vols = [float(x) for x in kl.get("vol", [])] if "vol" in kl else []

            bars_weeks = len(times)
            if args.min_history_weeks > 0 and bars_weeks < args.min_history_weeks:
                continue
            if times:
                first_ts = times[0]
                age_weeks = (now_ts - first_ts) / WEEK_SECS
                if args.min_age_weeks > 0 and age_weeks < args.min_age_weeks:
                    continue
            else:
                continue
            if len(highs) < 5 or len(lows) < 5:
                continue

            widx = last_complete_week_index(times, now_ts)
            weekly_amount = float("nan")
            weekly_vol = float("nan")
            if widx is not None:
                if amnts and abs(widx) <= len(amnts):
                    weekly_amount = amnts[widx]
                if vols and abs(widx) <= len(vols):
                    weekly_vol = vols[widx]

            if args.min_weekly_amount > 0 and (math.isnan(weekly_amount) or weekly_amount < args.min_weekly_amount):
                continue
            if args.min_weekly_vol > 0 and (math.isnan(weekly_vol) or weekly_vol < args.min_weekly_vol):
                continue

            bid1 = liq[sym].get("bid1")
            ask1 = liq[sym].get("ask1")
            spread_abs: Optional[float] = None
            spread_bps: Optional[float] = None
            mid: Optional[float] = None
            if args.max_spread_bps > 0:
                if bid1 is None or ask1 is None or bid1 <= 0 or ask1 <= 0 or ask1 < bid1:
                    continue
                mid = (bid1 + ask1) / 2.0
                spread_abs = ask1 - bid1
                spread_bps = (spread_abs / mid) * 10000.0
                if spread_bps > args.max_spread_bps:
                    continue
            else:
                if bid1 is not None and ask1 is not None and bid1 > 0 and ask1 > 0 and ask1 >= bid1:
                    mid = (bid1 + ask1) / 2.0
                    spread_abs = ask1 - bid1
                    spread_bps = (spread_abs / mid) * 10000.0

            hhv, llv = compute_hhv_llv(highs, lows, args.lookback_weeks)
            if math.isnan(hhv) or math.isnan(llv) or hhv <= llv:
                continue

            last_price = liq[sym]["lastPrice"] or (closes[-1] if closes else None)
            if last_price is None or last_price <= 0:
                continue

            pos = (last_price - llv) / (hhv - llv)
            pos_clipped = max(0.0, min(1.0, pos))
            prox_score = 1.0 - pos_clipped
            pct = pct_in_range(last_price, llv, hhv)
            dist_from_llv_pct = (last_price - llv) / llv * 100.0 if llv > 0 else math.nan

            last_ts = times[-1]
            last_wk = datetime.fromtimestamp(last_ts, tz=timezone.utc).strftime("%Y-%m-%d")

            amount24 = liq[sym]["amount24"]
            volume24 = liq[sym]["volume24"]
            holdvol = liq[sym]["holdVol"]

            illiquid_reasons = []
            if amount24 < args.min_amount24:
                illiquid_reasons.append(f"amount24<{args.min_amount24:g}")
            if volume24 < args.min_volume24:
                illiquid_reasons.append(f"volume24<{args.min_volume24:g}")
            if holdvol < args.min_holdvol:
                illiquid_reasons.append(f"holdVol<{args.min_holdvol:g}")
            is_illiquid = len(illiquid_reasons) > 0
            if args.exclude_illiquid and is_illiquid:
                continue

            # До последней закрытой недели
            if widx is None:
                closed_upto = len(closes) - 1
            else:
                closed_upto = len(closes) + widx
                if closed_upto < 0:
                    continue
            highs_c = highs[: closed_upto + 1]
            lows_c = lows[: closed_upto + 1]
            closes_c = closes[: closed_upto + 1]
            amnts_c = amnts[: closed_upto + 1] if amnts else []
            opens_c = opens[: closed_upto + 1] if opens else []

            # LLV26
            hhv26, llv26 = compute_hhv_llv(highs_c, lows_c, args.lookback_weeks2)
            dist_from_llv_pct_26 = (last_price - llv26) / llv26 * 100.0 if llv26 and llv26 > 0 else math.nan
            llv26_vs_llv52_bps = (
                (abs(llv26 - llv) / llv * 10000.0)
                if llv > 0 and llv26 and not math.isnan(llv26)
                else math.nan
            )
            llv26_approx_52 = (
                llv26_vs_llv52_bps <= args.llv26_close_bps
                if isinstance(llv26_vs_llv52_bps, (int, float)) and not math.isnan(llv26_vs_llv52_bps)
                else False
            )

            # ATRw
            atrw = atr_sma(highs_c, lows_c, closes_c, args.atrw_period, len(closes_c) - 1)
            atrw_pct = (atrw / last_price * 100.0) if atrw and last_price > 0 else math.nan

            # EMA 30/50/200, slope
            def last_ema(vals: Sequence[float], p: int) -> Tuple[Optional[float], Optional[float]]:
                if not vals:
                    return None, None
                es = ema_series(vals, p)
                if len(es) == 1:
                    return es[-1], None
                return es[-1], es[-2]

            ema30, ema30_prev = last_ema(closes_c, 30)
            ema50, _ = last_ema(closes_c, 50)
            ema200, _ = last_ema(closes_c, 200)
            ema30_slope = (ema30 - ema30_prev) if (ema30 is not None and ema30_prev is not None) else None
            trend_ok = (ema30 is not None and ema50 is not None and last_price >= ema30 >= ema50)

            # Weekly amount ratio vs median 26w
            wk_med = None
            wk_amount_ratio = None
            if amnts_c and len(amnts_c) >= 5 and not math.isnan(weekly_amount):
                import statistics

                take = amnts_c[-min(len(amnts_c), 26) :]
                if take:
                    wk_med = statistics.median(take)
                    if wk_med and wk_med > 0:
                        wk_amount_ratio = weekly_amount / wk_med

            # Touch cluster near LLV52
            touch_eps = args.touch_eps_bps / 10000.0
            wc = min(args.touch_weeks, len(lows_c))
            touch_count = 0
            if wc > 0 and llv > 0:
                for x in lows_c[-wc:]:
                    if x <= llv * (1.0 + touch_eps):
                        touch_count += 1
            cluster_ok = touch_count >= 2

            # Wickiness median 26w
            wickiness = None
            if opens_c and len(opens_c) >= 10:
                import statistics

                sample = []
                span = min(26, len(highs_c), len(lows_c), len(opens_c), len(closes_c))
                for i in range(-span, 0):
                    h, l, o, c = highs_c[i], lows_c[i], opens_c[i], closes_c[i]
                    rng = h - l
                    if rng <= 0:
                        continue
                    w = (abs(h - c) + abs(o - l)) / rng
                    sample.append(w)
                if sample:
                    wickiness = statistics.median(sample)

            # LR slope 13w
            lr_slope_13w = None
            if len(closes_c) >= 13:
                y = closes_c[-13:]
                n = len(y)
                xs = list(range(n))
                mx = sum(xs) / n
                my = sum(y) / n
                num = sum((x - mx) * (yy - my) for x, yy in zip(xs, y))
                den = sum((x - mx) ** 2 for x in xs) or 1e-9
                slope = num / den
                lr_slope_13w = (slope / last_price) * 100.0 if last_price > 0 else None

            # Distance to LLV52 как прокси стопа
            stop_dist_pct = (last_price - llv) / last_price * 100.0 if last_price > 0 else math.nan

            rows.append(
                {
                    "symbol": sym,
                    "last_price": round(last_price, 10),
                    "llv": round(llv, 10),
                    "hhv": round(hhv, 10),
                    "pos_in_channel": round(pos, 6),
                    "pct_in_range": round(pct, 4),
                    "dist_from_llv_pct": round(dist_from_llv_pct, 4),
                    "amount24": round(amount24, 4),
                    "volume24": round(volume24, 4),
                    "holdVol": round(holdvol, 4),
                    "weekly_amount": round(weekly_amount, 4) if not math.isnan(weekly_amount) else "",
                    "weekly_vol": round(weekly_vol, 4) if not math.isnan(weekly_vol) else "",
                    "bid1": round(bid1, 10) if bid1 is not None else "",
                    "ask1": round(ask1, 10) if ask1 is not None else "",
                    "spread_abs": round(spread_abs, 10) if spread_abs is not None else "",
                    "spread_bps": round(spread_bps, 4) if spread_bps is not None else "",
                    "mid_price": round(mid, 10) if mid is not None else "",
                    "illiquid": is_illiquid,
                    "illiquid_reasons": ",".join(illiquid_reasons),
                    "bars_weeks": bars_weeks,
                    "age_weeks": round(age_weeks, 2),
                    "last_week_close_time_utc": last_wk,
                    "lookback_weeks": args.lookback_weeks,
                    # NEW columns
                    "llv_26": round(llv26, 10) if llv26 and not math.isnan(llv26) else "",
                    "dist_from_llv_pct_26": (
                        round(dist_from_llv_pct_26, 4)
                        if isinstance(dist_from_llv_pct_26, (int, float)) and not math.isnan(dist_from_llv_pct_26)
                        else ""
                    ),
                    "llv26_vs_llv52_bps": (
                        round(llv26_vs_llv52_bps, 1)
                        if isinstance(llv26_vs_llv52_bps, (int, float)) and not math.isnan(llv26_vs_llv52_bps)
                        else ""
                    ),
                    "llv26_approx_52": bool(llv26_approx_52),
                    "atrw": round(atrw, 8) if isinstance(atrw, (int, float)) and not math.isnan(atrw) else "",
                    "atrw_pct": round(atrw_pct, 4)
                    if isinstance(atrw_pct, (int, float)) and not math.isnan(atrw_pct)
                    else "",
                    "ema30": round(ema30, 10) if isinstance(ema30, (int, float)) else "",
                    "ema50": round(ema50, 10) if isinstance(ema50, (int, float)) else "",
                    "ema200": round(ema200, 10) if isinstance(ema200, (int, float)) else "",
                    "ema30_slope": round(ema30_slope, 10)
                    if isinstance(ema30_slope, (int, float))
                    else "",
                    "trend_ok": bool(trend_ok),
                    "wk_amount_ratio": round(wk_amount_ratio, 4)
                    if isinstance(wk_amount_ratio, (int, float)) and not math.isnan(wk_amount_ratio)
                    else "",
                    "touch_count": touch_count,
                    "cluster_ok": bool(cluster_ok),
                    "wickiness": round(wickiness, 6)
                    if isinstance(wickiness, (int, float)) and not math.isnan(wickiness)
                    else "",
                    "lr_slope_13w": round(lr_slope_13w, 6)
                    if isinstance(lr_slope_13w, (int, float)) and not math.isnan(lr_slope_13w)
                    else "",
                    "stop_dist_pct": round(stop_dist_pct, 4)
                    if isinstance(stop_dist_pct, (int, float)) and not math.isnan(stop_dist_pct)
                    else "",
                    # Badges (эмодзи)
                    "badge_llv_match": ("🧱" if llv26_approx_52 else ""),
                    "badge_trend": ("📈" if trend_ok else ""),
                    "badge_volume": (
                        "💰"
                        if (
                            isinstance(wk_amount_ratio, (int, float))
                            and not math.isnan(wk_amount_ratio)
                            and wk_amount_ratio >= args.wk_amount_min_ratio
                        )
                        else ""
                    ),
                    "badges": (
                        ("🧱" if llv26_approx_52 else "")
                        + ("📈" if trend_ok else "")
                        + (
                            "💰"
                            if (
                                isinstance(wk_amount_ratio, (int, float))
                                and not math.isnan(wk_amount_ratio)
                                and wk_amount_ratio >= args.wk_amount_min_ratio
                            )
                            else ""
                        )
                    ),
                    "_prox_score": prox_score,
                    "_liq_raw": amount24,
                    "_closes_series": closes_c[-26:],
                    "_atrw_pct": atrw_pct if isinstance(atrw_pct, (int, float)) else math.nan,
                    "_wk_amount_ratio": wk_amount_ratio if isinstance(wk_amount_ratio, (int, float)) else math.nan,
                    "_spread_bps": spread_bps if isinstance(spread_bps, (int, float)) else math.nan,
                    "_stop_dist_pct": stop_dist_pct if isinstance(stop_dist_pct, (int, float)) else math.nan,
                }
            )
        except Exception:
            continue

    if not rows:
        log("Нет данных: ослабьте фильтры (history/age/liquidity/weekly/spread) или проверьте сеть/VPN.")
        sys.exit(3)

    # Жёсткий фильтр по дистанции до недельного лоя (если задан)
    if args.max_dist_w_pct and args.max_dist_w_pct > 0:
        before_n = len(rows)
        rows = [
            r
            for r in rows
            if _is_number(r.get("dist_from_llv_pct")) and r["dist_from_llv_pct"] <= args.max_dist_w_pct
        ]
        log(f"Фильтр по dist_from_llv_pct ≤ {args.max_dist_w_pct:g}%: осталось {len(rows)} из {before_n}")
        if not rows:
            log("После фильтрации ничего не осталось.")
            sys.exit(4)

    # Старый score (для совместимости консоли)
    liq_norm_map = normalize_log([r["_liq_raw"] for r in rows])
    for i, r in enumerate(rows):
        liq_norm = liq_norm_map.get(i, 0.0)
        score = w_prox * r["_prox_score"] + w_liq * liq_norm
        r["liq_norm"] = round(liq_norm, 6)
        r["score"] = round(score, 6)

    # BTC корреляция
    btc_closes: Optional[List[float]] = None
    if args.btc_corr:
        btc_quote = args.quote.upper() if args.quote else "USDT"
        btc_symbol = f"BTC_{btc_quote}"
        btc_kl = kl_map.get(btc_symbol)
        if not btc_kl:
            log(f"⚠ BTC kline не найден для {btc_symbol}, корреляция недоступна")
        else:
            try:
                times = [int(x) for x in btc_kl.get("time", [])]
                closes = [float(x) for x in btc_kl.get("close", [])]
                widx = last_complete_week_index(times, now_ts)
                if widx is None:
                    closed_upto = len(closes) - 1
                else:
                    closed_upto = len(closes) + widx
                if closed_upto >= 0:
                    btc_closes = closes[: closed_upto + 1][-26:]
            except Exception:
                btc_closes = None
        for r in rows:
            corr = None
            if btc_closes and r.get("_closes_series"):
                corr = corr_pearson(r["_closes_series"], btc_closes)
            if corr is None:
                r["btc_corr_26w"] = ""
            else:
                r["btc_corr_26w"] = round(corr, 4)

    # score_v2 и статусы
    if args.score2:
        near_scores: List[float] = []
        vol_scores: List[float] = []
        trend_scores: List[float] = []
        for r in rows:
            dist = r.get("dist_from_llv_pct")
            near_score = 0.0
            if _is_number(dist) and args.near_ll_pct > 0:
                near_score = max(0.0, 1.0 - min(dist / args.near_ll_pct, 1.5))
            r["_near_score"] = near_score
            near_scores.append(near_score)

            ratio = r.get("_wk_amount_ratio")
            vol_score = 0.0
            if isinstance(ratio, (int, float)) and ratio > 0:
                vol_score = min(ratio / max(args.wk_amount_min_ratio, 1e-6), 2.0) / 2.0
            r["_vol_score"] = vol_score
            vol_scores.append(vol_score)

            trend_score = 1.0 if r.get("trend_ok") else 0.0
            ema_slope = _safe_float(r.get("ema30_slope"))
            if ema_slope and ema_slope > 0:
                trend_score = min(1.0, trend_score + 0.25)
            trend_scores.append(trend_score)
            r["_trend_score"] = trend_score

        near_norm = {i: v for i, v in enumerate(near_scores)}
        liq_norm2 = {i: rows[i].get("liq_norm", 0.0) for i in range(len(rows))}
        vol_norm = {i: v for i, v in enumerate(vol_scores)}
        trend_norm = {i: v for i, v in enumerate(trend_scores)}
        total_weight = max(1e-9, args.w2_near + args.w2_liq + args.w2_vol + args.w2_trend)

        for i, r in enumerate(rows):
            score_v2 = (
                args.w2_near * near_norm.get(i, 0.0)
                + args.w2_liq * liq_norm2.get(i, 0.0)
                + args.w2_vol * vol_norm.get(i, 0.0)
                + args.w2_trend * trend_norm.get(i, 0.0)
            ) / total_weight
            r["score_v2"] = round(score_v2, 6)

            dist = r.get("dist_from_llv_pct")
            spread = r.get("_spread_bps")
            atr_pct = r.get("_atrw_pct")
            wk_ratio = r.get("_wk_amount_ratio")
            stop_pct = r.get("_stop_dist_pct")
            corr_val = _safe_float(r.get("btc_corr_26w")) if args.btc_corr else None

            near_ok = _is_number(dist) and dist <= args.near_ll_pct
            spread_ok = _is_number(spread) and spread <= args.watch_spread_bps if args.watch_spread_bps else True
            watch = bool(near_ok and spread_ok)

            ready = bool(
                watch
                and _is_number(wk_ratio)
                and wk_ratio >= args.wk_amount_min_ratio
                and _is_number(atr_pct)
                and atr_pct <= args.atrw_pct_max
            )

            trigger = bool(
                ready
                and r.get("llv26_approx_52")
                and r.get("cluster_ok")
                and (_safe_float(stop_pct) is None or stop_pct <= args.stop_max_pct)
            )
            if trigger and args.ema_guard:
                trigger = trigger and r.get("trend_ok")
            if trigger and args.btc_corr and corr_val is not None:
                trigger = trigger and abs(corr_val) <= args.btc_corr_max

            status = ""
            if trigger:
                status = "TRIGGER"
            elif ready:
                status = "READY"
            elif watch:
                status = "WATCH"
            r["status"] = status

    if args.sort_by == "score":
        rows.sort(key=lambda r: (-r.get("score", 0.0), r.get("pos_in_channel", 0.0)))
    elif args.sort_by == "pos":
        rows.sort(key=lambda r: (r.get("pos_in_channel", 0.0), -r.get("amount24", 0.0)))
    elif args.sort_by == "amount24":
        rows.sort(key=lambda r: (-r.get("amount24", 0.0), r.get("pos_in_channel", 0.0)))
    elif args.sort_by == "dist_w":
        rows.sort(key=lambda r: (
            r.get("dist_from_llv_pct", float("inf")),
            -r.get("amount24", 0.0),
        ))

    # Консольный топ
    topN = max(0, args.top)
    log("\n=== Топ ближайших к недельным лоям (score: ближе к 1 — лучше) ===")
    header = (
        f"{'SYMBOL':<18} {'PRICE':>11} {'LLV':>11} {'HHV':>11} {'POS':>6} {'PROX':>6} "
        f"{'SPbps':>7} {'LIQn':>6} {'SCORE':>7} {'$24h':>12} {'VOL24':>12} {'W$':>10} {'WVOL':>10} {'OI':>10} {'LQ?':>4}"
    )
    log(header)
    log("-" * len(header))
    for r in rows[:topN]:
        spbps = r['spread_bps'] if r['spread_bps'] != "" else 0
        wamt = r['weekly_amount'] if r['weekly_amount'] != "" else 0
        wvol = r['weekly_vol'] if r['weekly_vol'] != "" else 0
        log(
            f"{r['symbol']:<18} {r['last_price']:>11.6f} {r['llv']:>11.6f} {r['hhv']:>11.6f} "
            f"{r['pos_in_channel']:>6.4f} {r['_prox_score']:>6.4f} {spbps:>7} {r['liq_norm']:>6.3f} {r['score']:>7.4f} "
            f"{r['amount24']:>12.0f} {r['volume24']:>12.0f} {wamt:>10} {wvol:>10} {r['holdVol']:>10.0f} {('Y' if r['illiquid'] else ''):>4}"
        )

    export_rows: List[Dict] = []
    export_map: Dict[str, Dict] = {}
    for r in rows:
        cleaned = {k: v for k, v in r.items() if not k.startswith("_")}
        export_rows.append(cleaned)
        export_map[r["symbol"]] = cleaned

    csv_path = args.csv
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        if export_rows:
            keys = list(export_rows[0].keys())
            extra = sorted(set().union(*[set(r.keys()) for r in export_rows]) - set(keys))
            fieldnames = keys + extra
        else:
            fieldnames = ["empty"]
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        if export_rows:
            writer.writerows(export_rows)
    log(f"CSV: {csv_path}")

    # XLSX: сортировка по близости к недельному лою (ASC)
    excel_rows_sorted = sorted(
        [dict(r) for r in export_rows],
        key=lambda r: (
            float("inf")
            if (r.get("dist_from_llv_pct") in ("", None))
            else r.get("dist_from_llv_pct", float("inf")),
            -r.get("amount24", 0),
        ),
    )
    for i, rr in enumerate(excel_rows_sorted, 1):
        rr["rank_dist_w"] = i

    # Month1 check
    month_added = False
    if args.month_check and rows:
        symbols_week_passed = [r["symbol"] for r in rows]
        log(f"→ Klines Month1 (async) для {len(symbols_week_passed)} тикеров…")
        try:
            kl_m_map = asyncio.run(
                gather_klines_month(
                    symbols_week_passed,
                    months=args.lookback_months,
                    concurrency=min(args.concurrency, 8),
                )
            )
        except Exception:
            kl_m_map = {}
        if not kl_m_map:
            log("⚠ async(Month1) недоступен — перехожу на синхронную загрузку kline Month1.")
            kl_m_map = {}
            for i, sym in enumerate(symbols_week_passed, 1):
                try:
                    end_ts = int(time.time())
                    start_ts = end_ts - (args.lookback_months + 14) * MONTH_SECS
                    data = jget(
                        f"/api/v1/contract/kline/{sym}",
                        params={"interval": "Month1", "start": start_ts, "end": end_ts},
                    )
                    if isinstance(data, dict) and data.get("time"):
                        kl_m_map[sym] = data
                except Exception:
                    pass
                if i % 20 == 0:
                    log(f"…получено Month1 {i}/{len(symbols_week_passed)}")
        for r in rows:
            sym = r["symbol"]
            klm = kl_m_map.get(sym)
            if not klm:
                continue
            try:
                mh = [float(x) for x in klm.get("high", [])]
                ml = [float(x) for x in klm.get("low", [])]
                if len(mh) < 2 or len(ml) < 2:
                    continue
                m_hhv, m_llv = compute_hhv_llv(mh, ml, args.lookback_months)
                if math.isnan(m_hhv) or math.isnan(m_llv) or m_hhv <= m_llv:
                    continue
                wm_diff_bps = (abs((m_llv - r["llv"]) / r["llv"]) * 10000.0) if r["llv"] > 0 else float("nan")
                wm_close = not math.isnan(wm_diff_bps) and wm_diff_bps <= args.wm_llv_close_bps
                r["llv_m"] = round(m_llv, 10)
                r["hhv_m"] = round(m_hhv, 10)
                r["wm_llv_diff_bps"] = round(wm_diff_bps, 1) if not math.isnan(wm_diff_bps) else ""
                r["wm_llv_close"] = bool(wm_close)
                r["badge_wm"] = "✓WM" if wm_close else ""
                r["badges"] = (r.get("badges", "") or "") + ("✓WM" if wm_close else "")
                expo = export_map.get(sym)
                if expo is not None:
                    expo.update(
                        {
                            "llv_m": round(m_llv, 10),
                            "hhv_m": round(m_hhv, 10),
                            "wm_llv_diff_bps": round(wm_diff_bps, 1) if not math.isnan(wm_diff_bps) else "",
                            "wm_llv_close": bool(wm_close),
                            "badge_wm": "✓WM" if wm_close else "",
                            "badges": expo.get("badges", "") + ("✓WM" if wm_close else ""),
                        }
                    )
                month_added = True
            except Exception:
                continue
        if month_added:
            for rr in excel_rows_sorted:
                sym = rr.get("symbol")
                if not sym or sym not in export_map:
                    continue
                expo = export_map[sym]
                for key in ("llv_m", "hhv_m", "wm_llv_diff_bps", "wm_llv_close", "badge_wm", "badges"):
                    if key in expo:
                        rr[key] = expo[key]

    near_rows = [
        rr
        for rr in excel_rows_sorted
        if _is_number(rr.get("dist_from_llv_pct")) and rr.get("dist_from_llv_pct") <= args.near_ll_pct
    ]

    legend_rows = [
        {
            "Badge": "🧱",
            "Название": "LLV26≈LLV52",
            "Условие": f"|LLV26−LLV52| ≤ {args.llv26_close_bps:.0f} б.п.",
            "Смысл": "Короткое окно подтверждает годовой минимум",
        },
        {
            "Badge": "📈",
            "Название": "trend_ok",
            "Условие": "last_price ≥ EMA30 ≥ EMA50",
            "Смысл": "Минимальная тренд-защита (не ловим нож)",
        },
        {
            "Badge": "💰",
            "Название": "wk_amount_ratio",
            "Условие": f"weekly_amount ≥ {args.wk_amount_min_ratio:g} × медиана 26w",
            "Смысл": "Объём-подтверждение уровня",
        },
        {
            "Badge": "✓WM",
            "Название": "wm_llv_close",
            "Условие": f"|LLV_month−LLV_week| ≤ {args.wm_llv_close_bps:.0f} б.п.",
            "Смысл": "Месячный и недельный лои рядом",
        },
        {
            "Badge": "WATCH",
            "Название": "статус",
            "Условие": f"dist_from_llv_pct ≤ {args.near_ll_pct:g}% и spread_bps ≤ {args.watch_spread_bps:g}",
            "Смысл": "Близко к лою, ликвидно",
        },
        {
            "Badge": "READY",
            "Название": "статус",
            "Условие": f"WATCH + wk_amount_ratio ≥ {args.wk_amount_min_ratio:g} и atrw_pct ≤ {args.atrw_pct_max:g}%",
            "Смысл": "Есть объём и умеренная вола",
        },
        {
            "Badge": "TRIGGER",
            "Название": "статус",
            "Условие": "READY + llv26≈52 + cluster_ok"
            + (" + trend_ok" if args.ema_guard else "")
            + (f" + |corr| ≤ {args.btc_corr_max:g}" if args.btc_corr else "")
            + f" + stop_dist_pct ≤ {args.stop_max_pct:g}%",
            "Смысл": "Строгие условия входа",
        },
        {
            "Badge": "Near_LL",
            "Название": "лист",
            "Условие": f"dist_from_llv_pct ≤ {args.near_ll_pct:g}%",
            "Смысл": "Текущая цена действительно рядом с недельным лоем",
        },
    ]

    settings_rows = [
        {"Параметр": "lookback_weeks", "Значение": args.lookback_weeks},
        {"Параметр": "lookback_weeks2", "Значение": args.lookback_weeks2},
        {"Параметр": "lookback_months", "Значение": args.lookback_months},
        {"Параметр": "near_ll_pct", "Значение": f"{args.near_ll_pct:g}%"},
        {"Параметр": "watch_spread_bps", "Значение": f"{args.watch_spread_bps:g} б.п."},
        {"Параметр": "wk_amount_min_ratio", "Значение": args.wk_amount_min_ratio},
        {"Параметр": "atrw_pct_max", "Значение": f"{args.atrw_pct_max:g}%"},
        {"Параметр": "ema_guard", "Значение": str(bool(args.ema_guard))},
        {"Параметр": "btc_corr", "Значение": str(bool(args.btc_corr))},
        {"Параметр": "btc_corr_max", "Значение": args.btc_corr_max},
        {"Параметр": "stop_max_pct", "Значение": f"{args.stop_max_pct:g}%"},
        {"Параметр": "max_dist_w_pct", "Значение": f"{args.max_dist_w_pct:g}%"},
    ]

    sheets = {
        "All": excel_rows_sorted,
        "Near_LL": near_rows,
        "Legend": legend_rows,
        "Settings": settings_rows,
    }

    xlsx_path = args.xlsx
    if save_xlsx_multi(sheets, xlsx_path):
        log(f"XLSX: {xlsx_path}")
    else:
        log("⚠ Не удалось сохранить XLSX (нет openpyxl/pandas?)")

    if args.gsheet and export_rows:
        if upload_to_gsheet(export_rows, args.gsheet):
            log(f"Google Sheets: {args.gsheet}")
        else:
            log("⚠ Не удалось обновить Google Sheets")

    if args.telegram and export_rows:
        top_rows = export_rows[: min(len(export_rows), args.top or 10)]
        lines = [
            "<b>MEXC Futures — LLV52 Radar</b>",
            f"Всего: {len(export_rows)}",
        ]
        for r in top_rows:
            badges = r.get("badges", "")
            lines.append(
                f"{r['symbol']}: price={r['last_price']} dist={r['dist_from_llv_pct']}% {badges}"
            )
        if send_telegram("\n".join(lines)):
            log("Telegram: отправлено")
        else:
            log("⚠ Telegram: ошибка отправки")


if __name__ == "__main__":
    main()
